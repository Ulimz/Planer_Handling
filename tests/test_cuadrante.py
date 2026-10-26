from datetime import date

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from planer_handling.desde_cuadrante import generar_parrilla_cuadrante
from planer_handling.importar.cuadrante import leer_cuadrante, tramos
from planer_handling.modelos import Vuelo


def test_tramos_partidos_y_cruce_de_medianoche():
    assert tramos("05:00\n12:45") == [(300, 765)]
    assert tramos("15:00_x000D_\n01:00") == [(900, 1500)]
    assert tramos("04:45\n07:10\n\n08:50\n11:25") == [(285, 430), (530, 685)]
    assert tramos("") == []


def _cuadrante(ruta):
    wb = Workbook()
    ws = wb.active
    ws.append(["Depart", "Empleado"] + [f"{d}.2\n{n}" for d, n in zip(range(1, 8), "sdlmxjv")])
    filas = [("TL", "UNO, A", "05:00\n12:45"), ("PB*/JAR", "DOS, B", "05:00\n12:45"),
             ("PCP", "TRES, C", "05:00\n10:00"), ("PCP", "CUATRO, D", "04:45\n07:10\n\n08:50\n11:25"),
             ("PB*/JAR", "CINCO, E", "05:00\n12:45"), ("JAR", "SEIS, F", "05:00\n13:00"),
             ("RCO", "SIETE, G", "05:15\n15:15")]
    for funcion, nombre, horario in filas:
        ws.append([funcion, nombre] + [horario] * 3 + [None] * 4)
    ws["C3"].fill = PatternFill("solid", fgColor="FFFFFF00")
    wb.save(ruta)


def test_leer_cuadrante(tmp_path):
    ruta = tmp_path / "turnos.xlsx"
    _cuadrante(ruta)
    filas = leer_cuadrante(ruta, 2025)
    assert [f.funcion for f in filas][:2] == ["TL", "PB*/JAR"]
    assert set(filas[0].jornadas) == {date(2025, 2, 1), date(2025, 2, 2), date(2025, 2, 3)}
    assert filas[1].jornadas[date(2025, 2, 1)].marca == "FFFFFF00"
    assert len(filas[3].jornadas[date(2025, 2, 1)].tramos) == 2


def test_parrilla_desde_cuadrante(tmp_path, conf):
    ruta = tmp_path / "turnos.xlsx"
    _cuadrante(ruta)
    dia = date(2025, 2, 1)
    vuelos = [Vuelo(dia, vuelo_out="FR1", destino="STN", std=360),
              Vuelo(dia, vuelo_in="FR2", origen="DUB", sta=500, vuelo_out="FR3", destino="DUB", std=525, puesto="R05")]
    p = generar_parrilla_cuadrante(dia, vuelos, leer_cuadrante(ruta, 2025), conf, segundos=5)
    assert len(p.reparto.equipos) == 1
    equipo = p.reparto.equipos[0]
    assert equipo.miembros["TL"].nombre == "UNO, A"
    assert equipo.miembros["PB/TOW"].contrato == "PB*/JAR"
    assert (equipo.inicio, equipo.fin) == (300, 765)
    assert all(t.id in p.asignadas for t in p.tareas)
    assert len(p.reparto.jardineras) == 1 and len(p.reparto.mandos) == 1


def test_color_de_vacaciones_deja_fuera_aunque_tenga_horas(tmp_path, conf):
    ruta = tmp_path / "turnos.xlsx"
    _cuadrante(ruta)  # la fila 3 (PB*/JAR, 1 de febrero) está en amarillo
    dia = date(2025, 2, 1)
    p = generar_parrilla_cuadrante(dia, [Vuelo(dia, vuelo_out="FR1", destino="STN", std=360)],
                                   leer_cuadrante(ruta, 2025), conf, segundos=5)
    ausencias = [i for i in p.incidencias if i.codigo == "AUSENCIA"]
    assert len(ausencias) == 1 and "Vacaciones" in ausencias[0].detalle
    assert p.reparto.equipos[0].miembros["PB/TOW"].contrato == "PB*/JAR"  # lo cubre la otra persona PB*/JAR


def test_cabecera_de_dia_sin_salto_de_linea(tmp_path):
    ruta = tmp_path / "turnos.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Depart.", "  Empleado"] + [f"{d}.10_x000D_{n}" for d, n in zip(range(1, 8), ["jue", "vie", "sáb", "dom", "lun", "mar", "mié"])])
    ws.append(["TL", "UNO, A"] + ["05:00\n12:45\n"] * 7)
    wb.save(ruta)
    filas = leer_cuadrante(ruta, 2026)
    assert min(filas[0].jornadas) == date(2026, 10, 1) and len(filas[0].jornadas) == 7


def test_cada_tl_abre_su_equipo(tmp_path, conf):
    ruta = tmp_path / "turnos.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Depart", "Empleado"] + [f"{d}.2\n{n}" for d, n in zip(range(1, 8), "sdlmxjv")])
    for funcion, nombre in [("TL", "UNO, A"), ("TL", "DOS, B"), ("PB", "TRES, C"), ("PCP", "CUATRO, D")]:
        ws.append([funcion, nombre] + ["05:00\n12:45"] * 7)
    wb.save(ruta)
    dia = date(2025, 2, 1)
    p = generar_parrilla_cuadrante(dia, [Vuelo(dia, vuelo_out="FR1", destino="STN", std=420)],
                                   leer_cuadrante(ruta, 2025), conf, segundos=5)
    assert len(p.reparto.equipos) == 2
    # Un solo vuelo: basta un equipo y el otro queda de refuerzo.
    assert sorted(e.seccion for e in p.reparto.equipos) == ["GRUPOS", "REFUERZO"]
