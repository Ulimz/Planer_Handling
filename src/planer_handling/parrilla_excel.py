"""Exporta la parrilla a Excel: equipos, vuelos, jardineras, resto del personal, incidencias y Gantt."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelos import Equipo, hhmm
from .planificador import Parrilla

CABECERA = PatternFill("solid", fgColor="1F4E78")
VACIO = PatternFill("solid", fgColor="F8CBAD")
PENDIENTE = PatternFill("solid", fgColor="F4B084")
AVISO = PatternFill("solid", fgColor="FFE699")
COLORES_GANTT = ["9BC2E6", "A9D08E", "FFD966", "F4B084", "C9A0DC", "8EA9DB"]


def _tabla(ws, cabeceras: list[str], filas: list[list]) -> None:
    ws.append(cabeceras)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = CABECERA
    for fila in filas:
        ws.append(fila)
    for i, cab in enumerate(cabeceras, 1):
        ancho = max([len(str(cab))] + [len(str(f[i - 1])) for f in filas if i - 1 < len(f)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 8), 60)
    ws.freeze_panes = "A2"


def _nombre(e) -> str:
    return e.id if e else "VACANTE"


def exportar(parrilla: Parrilla, ruta: str | Path, puestos: list[str]) -> Path:
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    r = parrilla.reparto
    tareas = {t.id: t for t in parrilla.tareas + parrilla.tareas_jardinera}
    por_recurso: dict[str, list] = {}
    for tid, codigo in parrilla.asignadas.items():
        por_recurso.setdefault(codigo, []).append(tareas[tid])

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    sin_asignar = [t for t in parrilla.tareas if t.id not in parrilla.asignadas]
    resumen = [
        ["Día", parrilla.dia.strftime("%d/%m/%Y")],
        ["Vuelos (filas de programación)", len(parrilla.vuelos)],
        ["Tareas de plataforma", len(parrilla.tareas)],
        ["Tareas sin equipo", len(sin_asignar)],
        ["Tareas de jardinera", len(parrilla.tareas_jardinera)],
        ["Equipos de plataforma", len(r.equipos)],
        ["Equipos incompletos", sum(1 for e in r.equipos if not e.completo)],
        ["Jardineras", len(r.jardineras)],
        ["Personal en reserva", len(r.reserva)],
        ["Incidencias pendientes", sum(1 for i in parrilla.incidencias if i.gravedad == "PENDIENTE")],
        ["Avisos", sum(1 for i in parrilla.incidencias if i.gravedad == "AVISO")],
    ]
    for turno in sorted({e.turno for e in r.equipos}):
        resumen.append([f"Equipos turno {turno}", sum(1 for e in r.equipos if e.turno == turno)])
    _tabla(ws, ["Dato", "Valor"], resumen)

    ws = wb.create_sheet("Equipos")
    filas = []
    for e in r.equipos:
        vuelos = sorted(por_recurso.get(e.codigo, []), key=lambda t: t.inicio)
        filas.append([e.codigo, e.turno, f"{hhmm(e.inicio)}-{hhmm(e.fin)}"]
                     + [_nombre(e.miembros[p]) for p in puestos]
                     + [len(vuelos), ", ".join(t.vuelo.codigo for t in vuelos)])
    _tabla(ws, ["Equipo", "Turno", "Horario"] + puestos + ["Nº vuelos", "Vuelos"], filas)
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            if celda.value == "VACANTE":
                celda.fill = VACIO

    ws = wb.create_sheet("Vuelos")
    filas = []
    for t in sorted(parrilla.tareas, key=lambda t: t.inicio):
        v = t.vuelo
        filas.append([v.vuelo_in, v.origen, hhmm(v.sta) if v.sta is not None else "",
                      v.vuelo_out, v.destino, hhmm(v.std) if v.std is not None else "",
                      v.tipo_avion, v.puesto, t.tipo, f"{hhmm(t.inicio)}-{hhmm(t.fin)}",
                      parrilla.asignadas.get(t.id, "SIN EQUIPO")])
    _tabla(ws, ["Vuelo in", "Origen", "STA", "Vuelo out", "Destino", "STD", "Avión", "Puesto",
                "Servicio", "Ventana", "Equipo"], filas)

    ws = wb.create_sheet("Jardineras")
    filas = []
    for t in sorted(parrilla.tareas_jardinera, key=lambda t: t.inicio):
        recurso = parrilla.asignadas.get(t.id, "SIN JARDINERA")
        conductor = next((j.miembros["Conductor"] for j in r.jardineras if j.codigo == recurso), None)
        filas.append([t.vuelo.codigo, t.tipo.replace("JARDINERA_", "").lower(),
                      f"{hhmm(t.inicio)}-{hhmm(t.fin)}", recurso, _nombre(conductor) if conductor or recurso != "SIN JARDINERA" else ""])
    _tabla(ws, ["Vuelo", "Servicio", "Ventana", "Jardinera", "Conductor"], filas)

    ws = wb.create_sheet("Personal")
    filas = []
    for grupo in (r.mandos, r.hipodromo, r.parciales):
        for e in grupo:
            for puesto, emp in e.miembros.items():
                filas.append([e.seccion, e.codigo, e.turno, f"{hhmm(e.inicio)}-{hhmm(e.fin)}", puesto, _nombre(emp)])
    for emp in r.reserva:
        filas.append(["RESERVA", "", "", "", emp.categoria, emp.id])
    _tabla(ws, ["Sección", "Código", "Turno", "Horario", "Puesto", "Empleado"], filas)

    ws = wb.create_sheet("Incidencias")
    orden = {"PENDIENTE": 0, "AVISO": 1}
    filas = [[i.gravedad, i.codigo, i.detalle, i.referencia]
             for i in sorted(parrilla.incidencias, key=lambda i: (orden.get(i.gravedad, 2), i.codigo))]
    _tabla(ws, ["Gravedad", "Código", "Detalle", "Referencia"], filas)
    for fila in ws.iter_rows(min_row=2):
        fila[0].fill = PENDIENTE if fila[0].value == "PENDIENTE" else AVISO

    _gantt(wb.create_sheet("Gantt"), r.equipos + r.jardineras, por_recurso)
    wb.save(ruta)
    return ruta


def _gantt(ws, recursos: list[Equipo], por_recurso: dict[str, list], paso: int = 15) -> None:
    if not recursos:
        return
    inicio = min(r.inicio for r in recursos) // 60 * 60
    fin = max(r.fin for r in recursos)
    franjas = list(range(inicio, fin, paso))
    ws.append(["Recurso"] + [hhmm(m) if m % 60 == 0 else "" for m in franjas])
    ws.column_dimensions["A"].width = 10
    for i in range(2, len(franjas) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 3.2
    for celda in ws[1]:
        celda.font = Font(bold=True, size=8)
    turno_fill = PatternFill("solid", fgColor="EDEDED")
    for n, r in enumerate(recursos, 2):
        ws.cell(n, 1, r.codigo).font = Font(bold=True)
        for j, m in enumerate(franjas, 2):
            if r.inicio <= m < r.fin:
                ws.cell(n, j).fill = turno_fill
        for k, t in enumerate(sorted(por_recurso.get(r.codigo, []), key=lambda t: t.inicio)):
            color = PatternFill("solid", fgColor=COLORES_GANTT[k % len(COLORES_GANTT)])
            primera = True
            for j, m in enumerate(franjas, 2):
                if t.inicio < m + paso and m < t.fin:
                    celda = ws.cell(n, j)
                    celda.fill = color
                    if primera:
                        celda.value = t.vuelo.vuelo_out or t.vuelo.vuelo_in
                        celda.font = Font(size=7)
                        celda.alignment = Alignment(horizontal="left")
                        primera = False
    ws.freeze_panes = "B2"
