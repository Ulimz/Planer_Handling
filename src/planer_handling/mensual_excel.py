"""Excel del cuadrante mensual: mismo formato que el cuadrante de la empresa (se puede volver a
leer con el modo --cuadrante) más hojas de cobertura y de horas por persona."""
from __future__ import annotations

import calendar
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .mensual import INICIO_MAS_TEMPRANO, PASO, Cambio, CuadranteMes, texto_jornada
from .modelos import hhmm

DIAS_SEMANA = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]
NEGRITA = Font(bold=True)
CABECERA = PatternFill("solid", fgColor="FFD9E1F2")
FALTA = PatternFill("solid", fgColor="FFF8CBAD")
AUSENTE = PatternFill("solid", fgColor="FFD9D9D9")


def exportar_mes(cm: CuadranteMes, ruta, con_nombres: bool = False, colores: dict[str, str] | None = None,
                 cambios: list[Cambio] | None = None) -> None:
    """`colores`: motivo de ausencia -> color de celda (los de la empresa), para pintar cada ausencia.
    `cambios`: si es una reorganización, la lista de cambios va en su propia hoja."""
    wb = Workbook()
    _hoja_cuadrante(wb.active, cm, con_nombres, colores or {})
    if cambios is not None:
        _hoja_cambios(wb.create_sheet("Cambios", 0), cambios)
    _hoja_cobertura(wb.create_sheet("Cobertura"), cm)
    _hoja_horas(wb.create_sheet("Horas"), cm)
    _hoja_curva(wb.create_sheet("Curva"), cm)
    wb.save(ruta)


def _hoja_cuadrante(ws, cm: CuadranteMes, con_nombres: bool, colores: dict[str, str]) -> None:
    ws.title = "Cuadrante"
    dias = [d.dia for d in cm.dias]
    ws.append(["Depart.", "Empleado", "Patrón"] + [f"{d.day}.{d.month}\n{DIAS_SEMANA[d.weekday()]}" for d in dias])
    por_persona = defaultdict(dict)
    for dp in cm.dias:
        for j in dp.jornadas:
            por_persona[j.perfil.empleado.id][dp.dia] = j
    for p in cm.perfiles:
        e = p.empleado
        nombre = e.nombre if con_nombres else e.id
        fila = [p.funcion, nombre, p.patron]
        for dp in cm.dias:
            j = por_persona[e.id].get(dp.dia)
            fila.append(texto_jornada(j.tramos) if j else dp.ausentes.get(e.id))
        ws.append(fila)
        for dp, celda in zip(cm.dias, ws[ws.max_row][3:]):
            if e.id in dp.ausentes:
                color = colores.get(dp.ausentes[e.id])
                celda.fill = PatternFill("solid", fgColor=color) if color else AUSENTE
                if color == "FF000000":
                    celda.font = Font(color="FFFFFFFF")
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
        celda.alignment = Alignment(wrap_text=True, horizontal="center")
    for col in range(4, 4 + len(dias)):
        ws.column_dimensions[get_column_letter(col)].width = 7
    for fila in ws.iter_rows(min_row=2, min_col=4):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, horizontal="center")
    ws.freeze_panes = "D2"


def _hoja_cobertura(ws, cm: CuadranteMes) -> None:
    ws.append(["Día", "Vuelos", "Pico de equipos", "Personas", "Horas planificadas",
               "Horas sin cubrir TL", "Horas sin cubrir PB/TOW", "Horas sin cubrir operarios"])
    for dp in cm.dias:
        fila = [dp.dia, dp.vuelos, max(dp.demanda["TL"]), len(dp.jornadas),
                round(sum(j.minutos for j in dp.jornadas) / 60, 1)]
        fila += [round(dp.deficit(r) / 60, 1) for r in ("TL", "PB/TOW", "Operario")]
        ws.append(fila)
        if any(v > 0 for v in fila[5:]):
            for celda in ws[ws.max_row][5:]:
                celda.fill = FALTA
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
    ws.column_dimensions["A"].width = 12


def _hoja_horas(ws, cm: CuadranteMes) -> None:
    ws.append(["Empleado", "Función", "Patrón", "Contrato", "Días mes anterior", "Horas mes anterior",
               "Días planificados", "Horas planificadas", "Horas convenio", "Aviso"])
    dias, minutos = defaultdict(int), defaultdict(int)
    for dp in cm.dias:
        for j in dp.jornadas:
            dias[j.perfil.empleado.id] += 1
            minutos[j.perfil.empleado.id] += j.minutos
    # Los límites del convenio son por mes entero: si se planifica solo una parte, en proporción.
    proporcion = 0.0
    if cm.dias:
        d0 = cm.dias[0].dia
        proporcion = len(cm.dias) / calendar.monthrange(d0.year, d0.month)[1]
    for p in cm.perfiles:
        e = p.empleado
        horas = round(minutos[e.id] / 60, 1)
        limites, aviso = "", ""
        if p.horas_mes:
            minimo = p.horas_mes[0] * proporcion
            maximo = p.horas_mes[1] * proporcion if len(p.horas_mes) > 1 and p.horas_mes[1] else None
            limites = f"{minimo:.0f}-{maximo:.0f}" if maximo is not None else f"mín. {minimo:.0f}"
            if horas < minimo:
                aviso = "Por debajo del mínimo"
            elif maximo is not None and horas > maximo:
                aviso = "Por encima del máximo"
        ws.append([e.id, p.funcion, p.patron, p.contrato, p.dias_referencia, round(p.minutos_referencia / 60, 1),
                   dias[e.id], horas, limites, aviso])
        if aviso:
            ws.cell(ws.max_row, 10).fill = FALTA
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA


def _hoja_curva(ws, cm: CuadranteMes) -> None:
    """Demanda frente a cobertura cada 15 minutos, para hacer gráficos."""
    ws.append(["Día", "Hora", "Rol", "Necesarios", "Planificados"])
    for dp in cm.dias:
        for rol in ("TL", "PB/TOW", "Operario"):
            for k, (d, c) in enumerate(zip(dp.demanda[rol], dp.cobertura[rol])):
                if d or c:
                    ws.append([dp.dia, hhmm(INICIO_MAS_TEMPRANO + k * PASO), rol, d, c])
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA


def _hoja_cambios(ws, cambios: list[Cambio]) -> None:
    ws.append(["Día", "Empleado", "Función", "Contrato", "Antes", "Ahora", "Tipo", "Días de preaviso",
               "Horas nuevas", "Tipo de horas (convenio)"])
    for c in sorted(cambios, key=lambda c: (c.dia, c.tipo, c.empleado)):
        ws.append([c.dia, c.empleado, c.funcion, c.contrato, c.antes, c.despues, c.tipo, c.preaviso,
                   c.horas or None, c.tipo_horas])
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
    for col, ancho in zip("ABCDEFGHIJ", (12, 32, 10, 9, 24, 24, 22, 10, 10, 44)):
        ws.column_dimensions[col].width = ancho
