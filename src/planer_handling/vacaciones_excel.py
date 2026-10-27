"""Excel del reparto de vacaciones: asignación por persona, calendario y cupos."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .vacaciones import Solicitud, orden_de_eleccion, quincenas

NEGRITA = Font(bold=True)
CABECERA = PatternFill("solid", fgColor="FFD9E1F2")
VACACIONES = PatternFill("solid", fgColor="FFFFFF00")     # amarillo, como en el cuadrante
SIN_PREFERENCIA = PatternFill("solid", fgColor="FFF8CBAD")
LLENO = PatternFill("solid", fgColor="FFF8CBAD")


def _cabecera(ws) -> None:
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
        celda.alignment = Alignment(wrap_text=True, horizontal="center")


def exportar_vacaciones(solicitudes: list[Solicitud], anio: int, usado: dict[tuple[str, str], int],
                        cupo: dict[tuple[str, str], int], valor: dict[str, float], ruta) -> None:
    periodos = quincenas(anio)
    wb = Workbook()

    ws = wb.active
    ws.title = "Asignación"
    ws.append(["Orden", "Empleado", "Función", "Puntos antes", "Periodos", "Preferencias conseguidas",
               "Pidió", "Puntos después"])
    for n, s in enumerate(orden_de_eleccion(solicitudes), 1):
        ws.append([n, s.empleado, s.funcion, s.puntos, ", ".join(s.asignados),
                   f"{s.conseguidas} de {len(s.asignados)}", ", ".join(s.preferencias), s.puntos_despues(valor)])
        if s.conseguidas < len(s.asignados):
            ws.cell(ws.max_row, 6).fill = SIN_PREFERENCIA
    _cabecera(ws)
    for col, ancho in zip("ABCDEFGH", (7, 14, 10, 10, 18, 14, 40, 10)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("Calendario")
    ws.append(["Empleado", "Función"] + [p.codigo for p in periodos])
    for s in sorted(solicitudes, key=lambda s: (s.funcion, s.empleado)):
        ws.append([s.empleado, s.funcion] + ["V" if p.codigo in s.asignados else None for p in periodos])
        for p, celda in zip(periodos, ws[ws.max_row][2:]):
            if p.codigo in s.asignados:
                celda.fill = VACACIONES
                celda.alignment = Alignment(horizontal="center")
    _cabecera(ws)
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("Cupos")
    funciones = sorted({s.funcion for s in solicitudes})
    ws.append(["Periodo", "Desde", "Hasta", "Puntos"] + [f"{f} (usado / cupo)" for f in funciones])
    for p in periodos:
        ws.append([p.codigo, p.desde, p.hasta, valor[p.codigo]]
                  + [f"{usado.get((p.codigo, f), 0)} / {cupo[(p.codigo, f)]}" for f in funciones])
        for f, celda in zip(funciones, ws[ws.max_row][4:]):
            if usado.get((p.codigo, f), 0) >= cupo[(p.codigo, f)]:
                celda.fill = LLENO
    _cabecera(ws)
    wb.save(ruta)
