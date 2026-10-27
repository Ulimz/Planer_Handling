"""Parrilla de pasaje del día: quién factura cada vuelo, quién lo embarca y quién está en los
puestos fijos (oficina, lost and found…), con la gente de pasaje que trabaja ese día según el
cuadrante. Las cifras salen de `pasaje` en la configuración (ver mensual.demanda_pasaje)."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .configuracion import Configuracion
from .importar.cuadrante import FilaCuadrante
from .modelos import Vuelo, hhmm

PASO = 15                 # se reparte franja a franja de 15 minutos; así caben los relevos


@dataclass
class TareaPasaje:
    puesto: str                       # "Facturación", "Embarque", "Oficina"…
    vuelo: str
    inicio: int
    fin: int
    plazas: int
    agentes: dict[str, list[list[int]]] = field(default_factory=dict)   # agente -> tramos [desde, hasta]
    sin_cubrir: int = 0                                                 # minutos-persona sin nadie


def tareas_pasaje(vuelos: list[Vuelo], conf: Configuracion) -> list[TareaPasaje]:
    regla = conf.datos.get("pasaje", {})
    if not regla.get("activo") or not vuelos:
        return []
    fac, emb = regla.get("facturacion", {}), regla.get("embarque", {})
    tareas, momentos = [], []

    def en_bloques(puesto: str, vuelo: str, inicio: int, fin: int, plazas: int) -> None:
        tareas.append(TareaPasaje(puesto, vuelo, inicio, fin, plazas))

    for v in vuelos:
        if v.std is not None:
            codigo = v.vuelo_out or v.codigo
            en_bloques("Facturación", codigo, v.std - fac.get("abre_antes_std", 150),
                       v.std - fac.get("cierra_antes_std", 40), fac.get("personas", 0))
            en_bloques("Embarque", codigo, v.std - emb.get("desde_antes_std", 45),
                       v.std - emb.get("hasta_antes_std", 0), emb.get("personas", 0))
            momentos += [v.std - fac.get("abre_antes_std", 150), v.std]
        if v.sta is not None:
            momentos.append(v.sta + regla.get("despues_ultima_llegada", 60))
    inicio, fin = min(momentos) - regla.get("antes_primera_facturacion", 15), max(momentos)
    for puesto, personas in regla.get("puestos_fijos", {}).items():
        en_bloques(puesto, "", inicio, fin, personas)
    return sorted(tareas, key=lambda t: (t.inicio, t.puesto))


def agentes_del_dia(filas: list[FilaCuadrante], dia: date, conf: Configuracion, categoria: str,
                    ausentes: set[str] = frozenset()) -> dict[str, list[tuple[int, int]]]:
    """Nombre -> tramos de quien es de esa categoría y trabaja ese día (sin los ausentes)."""
    funciones = conf.datos.get("funciones", {})
    marcas = conf.datos.get("marcas", {})
    resultado = {}
    for f in filas:
        j = f.jornadas.get(dia)
        marca = marcas.get(j.marca) if j else None
        if (funciones.get(f.funcion, {}).get("categoria") == categoria and j and j.tramos
                and not (marca and not marca["disponible"]) and f.nombre not in ausentes):
            resultado[f.nombre] = j.tramos
    return resultado


def asignar(tareas: list[TareaPasaje], agentes: dict[str, list[tuple[int, int]]]) -> None:
    """Reparto franja a franja: quien ya está en una tarea sigue en ella mientras esté de turno; las
    plazas libres se dan a quien está de turno y libre, primero a quien seguirá de turno más tiempo
    (menos relevos) y con menos minutos asignados."""
    if not tareas:
        return
    carga = {a: 0 for a in agentes}
    inicio = min(t.inicio for t in tareas) // PASO * PASO
    fin = max(t.fin for t in tareas)
    en_tarea: dict[int, set[str]] = {i: set() for i in range(len(tareas))}

    def de_turno(a: str, m: int) -> int | None:
        """Hasta cuándo sigue de turno `a` si lo está en el minuto m."""
        return next((f for i, f in agentes[a] if i <= m and m + PASO <= f), None)

    for m in range(inicio, fin, PASO):
        activas = [i for i, t in enumerate(tareas) if t.inicio <= m < t.fin]
        ocupados: set[str] = set()
        for i in activas:                                   # los que siguen
            seguir = {a for a in en_tarea[i] if de_turno(a, m) is not None}
            en_tarea[i] = set(sorted(seguir)[:tareas[i].plazas])
            ocupados |= en_tarea[i]
        for i in sorted(activas, key=lambda i: tareas[i].fin):
            t = tareas[i]
            libres = [a for a in agentes if a not in ocupados and de_turno(a, m) is not None]
            libres.sort(key=lambda a: (-min(de_turno(a, m), t.fin), carga[a]))
            nuevos = libres[:t.plazas - len(en_tarea[i])]
            en_tarea[i] |= set(nuevos)
            ocupados |= set(nuevos)
            t.sin_cubrir += PASO * (t.plazas - len(en_tarea[i]))
            for a in en_tarea[i]:
                carga[a] += PASO
                tramos = t.agentes.setdefault(a, [])
                if tramos and tramos[-1][1] == m:
                    tramos[-1][1] = m + PASO
                else:
                    tramos.append([m, m + PASO])
        for i in en_tarea:
            if i not in activas:
                en_tarea[i] = set()


def anadir_hoja(ruta, tareas: list[TareaPasaje], responsables: dict[str, list[tuple[int, int]]]) -> None:
    """Añade la hoja "Pasaje" al Excel de la parrilla del día."""
    wb = load_workbook(ruta)
    ws = wb.create_sheet("Pasaje")
    negrita, cabecera = Font(bold=True), PatternFill("solid", fgColor="FFD9E1F2")
    falta = PatternFill("solid", fgColor="FFF8CBAD")
    ws.append(["Responsables de pasaje"] + [f"{n} ({' + '.join(f'{hhmm(i)}-{hhmm(f)}' for i, f in t)})"
                                            for n, t in sorted(responsables.items(), key=lambda x: x[1][0])])
    ws.append([])
    ws.append(["Desde", "Hasta", "Puesto", "Vuelo", "Personas", "Asignados", "Minutos sin cubrir"])
    for celda in ws[ws.max_row]:
        celda.font, celda.fill = negrita, cabecera
    ws["A1"].font = negrita
    for t in tareas:
        # Quien hace la tarea entera, solo el código; quien hace un relevo, con sus horas.
        asignados = ", ".join(
            a if len(x) == 1 and x[0][0] < t.inicio + PASO and x[0][1] >= t.fin else f"{a} ({' + '.join(f'{hhmm(i)}-{hhmm(f)}' for i, f in x)})"
            for a, x in sorted(t.agentes.items(), key=lambda kv: kv[1][0][0]))
        ws.append([hhmm(t.inicio), hhmm(t.fin), t.puesto, t.vuelo, t.plazas, asignados, t.sin_cubrir or None])
        if t.sin_cubrir:
            ws.cell(ws.max_row, 7).fill = falta
    for col, ancho in zip("ABCDEFG", (8, 8, 16, 10, 10, 90, 12)):
        ws.column_dimensions[col].width = ancho
    wb.save(ruta)
