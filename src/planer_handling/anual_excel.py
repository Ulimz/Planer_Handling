"""Excel del calendario de rotaciones de un año: una fila por persona y una columna por día."""
from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .anual import Rotacion, dias_del_anio
from .mensual_excel import CABECERA, DIAS_SEMANA, NEGRITA

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
COLORES = {"M": PatternFill("solid", fgColor="FFBDD7EE"), "T": PatternFill("solid", fgColor="FFF8CBAD"),
           "?": PatternFill("solid", fgColor="FFEDEDED")}
CENTRADO = Alignment(horizontal="center", vertical="center")


def exportar_anio(rotaciones: list[Rotacion], anio: int, ruta, codigos: dict[str, str] | None = None) -> None:
    """`codigos`: nombre -> código para no mostrar nombres; si no se pasa, salen los nombres."""
    dias = dias_del_anio(anio)
    orden = sorted(rotaciones, key=lambda r: (r.patron == "irregular", r.funcion, r.nombre))
    wb = Workbook()
    ws = wb.active
    ws.title = str(anio)
    ws.append(["Depart.", "Empleado", "Rotación"] + [f"{d.day}.{d.month}\n{DIAS_SEMANA[d.weekday()]}" for d in dias])
    for r in orden:
        nombre = codigos.get(r.nombre, r.nombre) if codigos is not None else r.nombre
        ws.append([r.funcion, nombre, r.patron if r.patron != "irregular" else "sin rotación fija"]
                  + [r.dia(d) for d in dias])
        for celda in ws[ws.max_row][3:]:
            if celda.value in COLORES:
                celda.fill = COLORES[celda.value]
            celda.alignment = CENTRADO
    for celda in ws[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
        celda.alignment = Alignment(wrap_text=True, horizontal="center")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16
    for col in range(4, 4 + len(dias)):
        ws.column_dimensions[get_column_letter(col)].width = 4.5
    ws.freeze_panes = "D2"

    resumen = wb.create_sheet("Días por mes")
    resumen.append(["Depart.", "Empleado", "Rotación"] + MESES + ["Total", "Mañanas", "Tardes"])
    for r in orden:
        if r.patron == "irregular":
            continue
        marcas = {d: r.dia(d) for d in dias}
        por_mes = Counter(d.month for d, m in marcas.items() if m)
        cuenta = Counter(marcas.values())
        nombre = codigos.get(r.nombre, r.nombre) if codigos is not None else r.nombre
        resumen.append([r.funcion, nombre, r.patron] + [por_mes[m] for m in range(1, 13)]
                       + [sum(por_mes.values()), cuenta["M"], cuenta["T"]])
    for celda in resumen[1]:
        celda.font, celda.fill = NEGRITA, CABECERA
    resumen.column_dimensions["B"].width = 36

    leyenda = wb.create_sheet("Leyenda")
    for texto, clave_color in [("M: trabaja de mañana", "M"), ("T: trabaja de tarde", "T"),
                               ("Vacío: libra", None), ("Sin rotación fija: se le llama según la demanda", "?")]:
        leyenda.append(["", texto])
        if clave_color:
            leyenda.cell(leyenda.max_row, 1).fill = COLORES[clave_color]
    leyenda.append([])
    leyenda.append(["", "La rotación continúa desde los últimos cuadrantes; no se reinicia en enero."])
    leyenda.append(["", "Las vacaciones no están incluidas: se reparten aparte según el convenio."])
    leyenda.column_dimensions["B"].width = 70
    wb.save(ruta)
