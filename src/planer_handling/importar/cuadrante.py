"""Lectura del cuadrante mensual de turnos.

Formato: una fila por persona con DEPART. (función) y EMPLEADO, y una columna por día con
cabeceras como "1.2\\nsáb" o fechas. Cada celda trae uno o dos tramos "HH:MM\\nHH:MM";
una celda vacía es un día sin trabajo. El color de la celda se guarda como marca para revisarla.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from .vuelos import clave

ALIAS_NOMBRE = ["EMPLEADO", "EMPLEADOS", "NOMBRE", "TRABAJADOR", "APELLIDOS Y NOMBRE", "NOMBRE Y APELLIDOS"]
ALIAS_FUNCION = ["DEPART", "DEPART.", "DEPARTAMENTO", "DPTO", "DPTO.", "FUNCION"]
HORA = re.compile(r"\b([01]?\d|2[0-3]):([0-5]\d)\b")


@dataclass
class Jornada:
    tramos: list[tuple[int, int]]       # minutos desde las 00:00; fin > inicio aunque cruce medianoche
    marca: str = ""                     # color de la celda (RGB) si lo tiene


@dataclass
class FilaCuadrante:
    funcion: str
    nombre: str
    fila: int
    jornadas: dict[date, Jornada] = field(default_factory=dict)
    marcas_vacias: dict[date, str] = field(default_factory=dict)   # días sin horario pero con color (baja…)


def _limpiar(valor) -> str:
    return "" if valor is None else str(valor).replace("_x000D_", "").strip()


def tramos(texto: str) -> list[tuple[int, int]]:
    horas = [int(h) * 60 + int(m) for h, m in HORA.findall(_limpiar(texto))]
    resultado = []
    for inicio, fin in zip(horas[0::2], horas[1::2]):
        resultado.append((inicio, fin if fin > inicio else fin + 1440))
    return resultado


def _dia_columna(valor, anio: int) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _limpiar(valor)
    if m := re.match(r"(\d{1,2})[./-](\d{1,2})(?!\d)", texto):
        try:
            return date(anio, int(m[2]), int(m[1]))
        except ValueError:
            return None
    return None


def leer_cuadrante(ruta: str | Path, anio: int) -> list[FilaCuadrante]:
    libro = load_workbook(ruta)
    for hoja in libro.worksheets:
        for fila_cab in range(1, 16):
            cabeceras = [clave(_limpiar(c.value)) for c in hoja[fila_cab]]
            col_nombre = next((i for i, c in enumerate(cabeceras, 1) if c in ALIAS_NOMBRE), None)
            col_funcion = next((i for i, c in enumerate(cabeceras, 1) if c in ALIAS_FUNCION), None)
            dias = {i: d for i, c in enumerate(hoja[fila_cab], 1) if (d := _dia_columna(c.value, anio))}
            if col_nombre and col_funcion and len(dias) >= 7:
                return _leer_filas(hoja, fila_cab, col_nombre, col_funcion, dias)
    raise ValueError("No se encuentra una tabla de cuadrante con EMPLEADO, DEPART. y columnas de días.")


def _leer_filas(hoja, fila_cab, col_nombre, col_funcion, dias) -> list[FilaCuadrante]:
    filas = []
    for r in range(fila_cab + 1, hoja.max_row + 1):
        nombre = _limpiar(hoja.cell(r, col_nombre).value)
        if not nombre:
            continue
        registro = FilaCuadrante(_limpiar(hoja.cell(r, col_funcion).value).upper(), nombre, r)
        for col, dia in dias.items():
            celda = hoja.cell(r, col)
            t = tramos(_limpiar(celda.value))
            marca = celda.fill.fgColor.rgb if celda.fill is not None and celda.fill.fill_type else ""
            marca = marca if isinstance(marca, str) and marca not in ("00000000", "FFFFFFFF") else ""
            if t:
                registro.jornadas[dia] = Jornada(t, marca)
            elif marca:
                registro.marcas_vacias[dia] = marca
        filas.append(registro)
    return filas
