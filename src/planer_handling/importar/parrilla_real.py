"""Lectura de los vuelos de una parrilla diaria ya hecha (una hoja por día, como "15-10").

Cabeceras en la fila 4: Orig. | N.Vuelo | Dest. | N.Vuelo | Matricula | Pk | ETA | STA | STD | ... | Grupo.
La llegada ocupa Orig. y las dos columnas de N.Vuelo (compañía y número); la salida, Dest. y las dos
siguientes. Las horas son locales. Devuelve también el grupo que hizo cada vuelo, para comparar.
"""
from __future__ import annotations

import re
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from ..modelos import Vuelo
from .vuelos import clave


def _minutos(valor) -> int | None:
    if isinstance(valor, (datetime, time)):
        return valor.hour * 60 + valor.minute
    m = re.match(r"\s*(\d{1,2}):(\d\d)", str(valor or ""))
    return int(m[1]) * 60 + int(m[2]) if m else None


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def leer_vuelos_parrilla(ruta: str | Path, hoja: str, dia: date) -> tuple[list[Vuelo], dict[str, str]]:
    libro = load_workbook(ruta, data_only=True, read_only=True)
    ws = libro[hoja]
    filas = list(ws.iter_rows(min_row=1, max_row=400, values_only=True))
    fila_cab = next((i for i, f in enumerate(filas[:15]) if "N.VUELO" in [clave(v) for v in f]), None)
    if fila_cab is None:
        raise ValueError(f"La hoja {hoja} no tiene la cabecera N.Vuelo de una parrilla diaria.")
    cab = [clave(v) for v in filas[fila_cab]]
    c_orig, c_dest = cab.index("ORIG."), cab.index("DEST.")
    c_pk, c_sta, c_std = cab.index("PK"), cab.index("STA"), cab.index("STD")
    c_grupo = cab.index("GRUPO") if "GRUPO" in cab else None

    vuelos, grupos = [], {}
    for n, f in enumerate(filas[fila_cab + 1:], fila_cab + 2):
        compania_in, compania_out = _texto(f[c_orig + 1]), _texto(f[c_dest + 1])
        if not (re.fullmatch(r"[A-Z0-9]{2}", compania_in) or re.fullmatch(r"[A-Z0-9]{2}", compania_out)):
            continue
        v = Vuelo(dia, fila_origen=f"{hoja}!{n}", tipo_avion="")
        if compania_in and _texto(f[c_orig + 2]):
            v.vuelo_in, v.origen, v.sta = compania_in + _texto(f[c_orig + 2]), _texto(f[c_orig]), _minutos(f[c_sta])
        if compania_out and _texto(f[c_dest + 2]):
            v.vuelo_out, v.destino, v.std = compania_out + _texto(f[c_dest + 2]), _texto(f[c_dest]), _minutos(f[c_std])
        if v.sta is None and v.std is None:
            continue
        if v.sta is not None and v.std is not None and v.std < v.sta:
            v.std += 1440
        v.puesto = _texto(f[c_pk])
        vuelos.append(v)
        if c_grupo is not None:
            grupos[v.codigo] = _texto(f[c_grupo])
    return vuelos, grupos
